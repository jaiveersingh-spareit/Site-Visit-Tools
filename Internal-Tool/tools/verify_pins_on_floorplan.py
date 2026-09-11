#!/usr/bin/env python3
"""
verify_pins_on_floorplan.py — overlay an exported visit's pin coordinates onto its
floor plan image, to visually confirm the exported X/Y data lines up with where pins
actually are, without needing to reopen the app.

Built during the MFS Luxembourg (08 Sep 2026) field-test investigation to settle
whether the app's stored pin coordinates were correct (they were — see BLOCKERS.md,
"Code dive" section). Reuse this on any future export instead of re-deriving the check.

Usage:
    python3 verify_pins_on_floorplan.py <Site_Visit_Data.xlsx> <floor_plan_image> <out.png> [floor_number]

Reads the "Floor Plan Pins" sheet (Floor, Type, Label, Location, X %, Y %, Notes) and
draws a red crosshair + label at each pin's stored position on the given floor plan
image. If the crosshairs land where you'd expect the real pins to be, the coordinate
data is trustworthy and any visual "doesn't match" complaint is a rendering/style issue,
not a data issue — check the export's pin-drawing code before doubting the numbers.

Requires: openpyxl, Pillow (pip install openpyxl pillow --break-system-packages)
"""
import sys
import openpyxl
from PIL import Image, ImageDraw


def load_pins(xlsx_path, floor_filter=None):
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    if "Floor Plan Pins" not in wb.sheetnames:
        raise SystemExit(
            f"No 'Floor Plan Pins' sheet in {xlsx_path}. "
            f"Sheets found: {wb.sheetnames}. "
            "(Older exports used 'Field Notes' with 'Pin X%'/'Pin Y%' columns instead — "
            "adapt this loader if you're checking one of those.)"
        )
    ws = wb["Floor Plan Pins"]
    rows = list(ws.iter_rows(values_only=True))
    header, data = rows[0], rows[1:]
    idx = {name: i for i, name in enumerate(header)}
    pins = []
    for r in data:
        floor = r[idx["Floor"]]
        if floor_filter is not None and str(floor) != str(floor_filter):
            continue
        x, y = r[idx["X %"]], r[idx["Y %"]]
        if x in (None, "Not positioned") or y in (None, "Not positioned"):
            continue
        pins.append({
            "floor": floor,
            "type": r[idx["Type"]],
            "label": r[idx["Label"]],
            "x": float(x),
            "y": float(y),
        })
    return pins


def overlay(image_path, pins, out_path):
    im = Image.open(image_path).convert("RGB")
    W, H = im.size
    draw = ImageDraw.Draw(im)
    for p in pins:
        px, py = p["x"] / 100 * W, p["y"] / 100 * H
        draw.ellipse([px - 6, py - 6, px + 6, py + 6], outline=(255, 0, 0), width=3)
        draw.line([px - 10, py, px + 10, py], fill=(255, 0, 0), width=2)
        draw.line([px, py - 10, px, py + 10], fill=(255, 0, 0), width=2)
        draw.text((px + 8, py - 14), f"{p['label']} ({p['type']})", fill=(255, 0, 0))
    im.save(out_path)
    print(f"Saved {out_path} — {len(pins)} pins overlaid at {W}x{H}")


if __name__ == "__main__":
    if len(sys.argv) < 4:
        print(__doc__)
        sys.exit(1)
    xlsx_path, image_path, out_path = sys.argv[1:4]
    floor_filter = sys.argv[4] if len(sys.argv) > 4 else None
    pins = load_pins(xlsx_path, floor_filter)
    if not pins:
        print("No positioned pins found for the given floor filter.")
        sys.exit(1)
    overlay(image_path, pins, out_path)
